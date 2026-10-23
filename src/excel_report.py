"""
Builds outputs/hedging_audit_report.xlsx.

Settlement and collar pricing are written as live Excel formulas, so a reviewer can
change a strike or the volatility and watch the workbook recompute. The Monte Carlo
results (10,000 scenarios per round) are written as values from stress_test.py,
since a simulation that size belongs in Python, not in cells.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from stress_test import N_SCENARIOS, PERIOD_YEARS, RISK_FREE, SEED, run

OUT = Path(__file__).resolve().parents[1] / "outputs" / "hedging_audit_report.xlsx"

FONT = "Arial"
BLUE, BLACK, GREEN = "0000FF", "000000", "008000"
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
TOTAL_FILL = PatternFill("solid", fgColor="EEF2F7")
THIN = Side(style="thin", color="C8CED8")
USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
MUSD = '$#,##0.00,,"M";($#,##0.00,,"M");"-"'
PCT = '0.0%;(0.0%);"-"'
NUM = '#,##0;(#,##0);"-"'


def f(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=FONT, color=color, bold=bold, size=size, italic=italic)


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = f(bold=True, size=14)
    ws["A2"] = sub
    ws["A2"].font = f(color="555555", italic=True)


def header(ws, row, labels, widths=None):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def put(ws, ref, value, fmt=None, color=BLACK, bold=False, fill=None):
    c = ws[ref]
    c.value = value
    c.font = f(color=color, bold=bold)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    c.border = Border(bottom=THIN)
    return c


def build():
    s, sigma, stress, pricing, sens = run()
    wb = Workbook()

    # ---------------- Assumptions ----------------
    wa = wb.active
    wa.title = "Assumptions"
    title(wa, "Assumptions", "Blue = input you can change. Black = formula. Yellow = key assumption.")
    header(wa, 4, ["Item", "Value", "Notes"], [34, 16, 90])
    rows = [
        ("Units produced per round", 100000, NUM, BLUE, None, "Game setup: producer sells 100,000 units each round."),
        ("Volatility per round (σ)", "=STDEV(Settlement!R6:R11)", PCT, BLACK, KEY_FILL,
         "Std dev of round-to-round log changes in the game's settled spot (Settlement col R). Feeds Collar Pricing."),
        ("Risk-free rate (annual)", RISK_FREE, PCT, BLUE, None, "Assumption, used only to discount option premiums."),
        ("Pricing horizon per round (years)", PERIOD_YEARS, "0.0000", BLUE, None, "One round treated as one month for discounting. σ above is already per round."),
        ("Annualized σ for Black-76", "=B6/SQRT(B8)", PCT, BLACK, None, "Per-round σ scaled so σ·√T equals one round of volatility."),
        ("Monte Carlo scenarios per round", N_SCENARIOS, NUM, BLUE, None, "Run in Python (src/stress_test.py). Change there and re-run."),
        ("Random seed", SEED, "0", BLUE, None, "Fixed so every run reproduces the same scenarios."),
        ("Forward price used for each round", "Swap price", None, BLUE, None,
         "The round's swap price is taken as the market's forward when I traded."),
    ]
    for i, (lab, val, fmt, col, fill, note) in enumerate(rows, 5):
        put(wa, f"A{i}", lab)
        put(wa, f"B{i}", val, fmt, col, fill=fill)
        put(wa, f"C{i}", note).font = f(color="555555")
    wa["A15"] = "Source for all positions and settled prices: my weekly decision memos, Oct 13 - Dec 2, 2025 (Google Drive, original version history)."
    wa["A15"].font = f(color="555555", italic=True)

    # ---------------- Settlement ----------------
    ws = wb.create_sheet("Settlement")
    title(ws, "Round-by-round settlement", "Recomputed from the memos with live formulas. Revenue = units x spot + swap P&L + collar P&L.")
    cols = ["Round", "Settled spot", "Swap qty", "Swap price", "Collar qty", "Put strike", "Call strike", "Float qty",
            "Unhedged revenue", "Swap P&L", "Collar P&L", "Hedged revenue", "Hedge P&L", "Memo revenue", "Matches memo",
            "Bankruptcy line", "Inv. hurdle", "Log change in spot"]
    header(ws, 4, cols, [8, 11, 10, 10, 10, 10, 10, 10, 15, 14, 14, 15, 14, 15, 11, 14, 14, 12])
    for i, (_, r) in enumerate(s.iterrows(), 5):
        put(ws, f"A{i}", int(r["round"]), "0", BLUE)
        put(ws, f"B{i}", float(r["spot"]), USD2, BLUE)
        put(ws, f"C{i}", int(r["swap_qty"]), NUM, BLUE)
        put(ws, f"D{i}", float(r["swap_price"]), USD2, BLUE)
        put(ws, f"E{i}", int(r["collar_qty"]), NUM, BLUE)
        put(ws, f"F{i}", None if r["collar_qty"] == 0 else float(r["put_strike"]), USD2, BLUE)
        put(ws, f"G{i}", None if r["collar_qty"] == 0 else float(r["call_strike"]), USD2, BLUE)
        put(ws, f"H{i}", f"=Assumptions!$B$5-C{i}-E{i}", NUM)
        put(ws, f"I{i}", f"=Assumptions!$B$5*B{i}", USD)
        put(ws, f"J{i}", f"=C{i}*(D{i}-B{i})", USD)
        put(ws, f"K{i}", f"=IF(E{i}=0,0,E{i}*(MAX(F{i}-B{i},0)-MAX(B{i}-G{i},0)))", USD)
        put(ws, f"L{i}", f"=I{i}+J{i}+K{i}", USD, bold=True)
        put(ws, f"M{i}", f"=L{i}-I{i}", USD)
        put(ws, f"N{i}", float(r["memo_reported_revenue"]), USD, BLUE)
        put(ws, f"O{i}", f'=IF(ABS(L{i}-N{i})<1,"Yes","CHECK")', None)
        put(ws, f"P{i}", None if r["bankruptcy_line"] != r["bankruptcy_line"] else float(r["bankruptcy_line"]), USD, BLUE)
        put(ws, f"Q{i}", None if r["investment_hurdle"] != r["investment_hurdle"] else float(r["investment_hurdle"]), USD, BLUE)
        put(ws, f"R{i}", "" if i == 5 else f"=LN(B{i}/B{i-1})", PCT)
    ws["P5"].comment = Comment("Not stated in my Round 1 memo.", "Karthik")
    ws["P10"].comment = Comment("Not stated in my Round 6 memo.", "Karthik")
    t = 12
    put(ws, f"A{t}", "Total", bold=True, fill=TOTAL_FILL)
    for c in "IJKLMN":
        put(ws, f"{c}{t}", f"=SUM({c}5:{c}11)", USD, bold=True, fill=TOTAL_FILL)
    put(ws, "A14", "Cost of hedging (unhedged - hedged)", bold=True)
    put(ws, "I14", "=I12-L12", USD, bold=True)
    put(ws, "A15", "Cost as % of unhedged revenue")
    put(ws, "I15", "=IF(I12=0,0,I14/I12)", PCT)
    put(ws, "A16", "Std dev of revenue, hedged")
    put(ws, "I16", "=STDEVP(L5:L11)", USD)
    put(ws, "A17", "Std dev of revenue, unhedged")
    put(ws, "I17", "=STDEVP(I5:I11)", USD)
    put(ws, "A18", "Volatility reduction from hedging")
    put(ws, "I18", "=IF(I17=0,0,1-I16/I17)", PCT, bold=True)
    put(ws, "A19", "Rounds below bankruptcy line, hedged")
    put(ws, "I19", '=SUMPRODUCT((P5:P11<>"")*(L5:L11<P5:P11))', "0")
    put(ws, "A20", "Rounds below bankruptcy line, unhedged")
    put(ws, "I20", '=SUMPRODUCT((P5:P11<>"")*(I5:I11<P5:P11))', "0")
    put(ws, "A21", "Rounds clearing the hurdle, hedged")
    put(ws, "I21", '=SUMPRODUCT((Q5:Q11<>"")*(L5:L11>=Q5:Q11))', "0")
    ws.freeze_panes = "B5"

    # ---------------- Collar Pricing ----------------
    wc = wb.create_sheet("Collar Pricing")
    title(wc, "Black-76 pricing of each collar I traded",
          "The game settled every collar at zero cost. These formulas price each leg at market and show the zero-cost call strike.")
    cc = ["Round", "Forward F", "Put strike", "Game call strike", "σ (annual)", "r", "T (yrs)",
          "d1 put", "d2 put", "Put premium", "d1 call", "d2 call", "Call premium", "Net premium /unit",
          "Collar qty", "Game's edge vs market", "Zero-cost call strike", "Call premium at zero-cost strike", "Solver check",
          "Collar delta /unit", "Collar gamma /unit"]
    header(wc, 4, cc, [8, 11, 11, 12, 11, 8, 9, 9, 9, 12, 9, 9, 12, 13, 11, 15, 14, 15, 11, 12, 12])
    for i, (_, p) in enumerate(pricing.iterrows(), 5):
        srow = 4 + int(p["round"])  # row of this round on Settlement
        put(wc, f"A{i}", int(p["round"]), "0", BLUE)
        put(wc, f"B{i}", f"=Settlement!D{srow}", USD2, GREEN)
        put(wc, f"C{i}", f"=Settlement!F{srow}", USD2, GREEN)
        put(wc, f"D{i}", f"=Settlement!G{srow}", USD2, GREEN)
        put(wc, f"E{i}", "=Assumptions!$B$9", PCT, GREEN)
        put(wc, f"F{i}", "=Assumptions!$B$7", PCT, GREEN)
        put(wc, f"G{i}", "=Assumptions!$B$8", "0.0000", GREEN)
        put(wc, f"H{i}", f"=(LN(B{i}/C{i})+0.5*E{i}^2*G{i})/(E{i}*SQRT(G{i}))", "0.0000")
        put(wc, f"I{i}", f"=H{i}-E{i}*SQRT(G{i})", "0.0000")
        put(wc, f"J{i}", f"=EXP(-F{i}*G{i})*(C{i}*NORMSDIST(-I{i})-B{i}*NORMSDIST(-H{i}))", USD2)
        put(wc, f"K{i}", f"=(LN(B{i}/D{i})+0.5*E{i}^2*G{i})/(E{i}*SQRT(G{i}))", "0.0000")
        put(wc, f"L{i}", f"=K{i}-E{i}*SQRT(G{i})", "0.0000")
        put(wc, f"M{i}", f"=EXP(-F{i}*G{i})*(B{i}*NORMSDIST(K{i})-D{i}*NORMSDIST(L{i}))", USD2)
        put(wc, f"N{i}", f"=J{i}-M{i}", USD2, bold=True)
        put(wc, f"O{i}", f"=Settlement!E{srow}", NUM, GREEN)
        put(wc, f"P{i}", f"=N{i}*O{i}", USD, bold=True)
        z = put(wc, f"Q{i}", round(float(p["zero_cost_call_strike"]), 4), USD2, BLUE)
        z.comment = Comment("Solved in Python (scipy brentq, src/pricing.py) at the volatility in Assumptions. "
                            "Column R reprices it in Excel so you can check it.", "Karthik")
        put(wc, f"R{i}", f"=EXP(-F{i}*G{i})*(B{i}*NORMSDIST((LN(B{i}/Q{i})+0.5*E{i}^2*G{i})/(E{i}*SQRT(G{i})))"
                         f"-Q{i}*NORMSDIST((LN(B{i}/Q{i})-0.5*E{i}^2*G{i})/(E{i}*SQRT(G{i}))))", USD2)
        put(wc, f"S{i}", f'=IF(ABS(R{i}-J{i})<0.01,"Matches put","Re-run solver")')
        put(wc, f"T{i}", f"=EXP(-F{i}*G{i})*((NORMSDIST(H{i})-1)-NORMSDIST(K{i}))", "0.000")
        put(wc, f"U{i}", f"=EXP(-F{i}*G{i})*(EXP(-H{i}^2/2)-EXP(-K{i}^2/2))/SQRT(2*PI())/(B{i}*E{i}*SQRT(G{i}))", "0.00000")
    n = 4 + len(pricing)
    put(wc, f"A{n+2}", "How to read it", bold=True)
    notes = [
        "Net premium /unit = put premium - call premium. Positive means the collar would have cost money at market prices.",
        "Game's edge vs market = what the game's free collar was worth to me. Positive = the game gave me a better deal than the market would.",
        "If the game's call strike sits above the zero-cost strike, I kept more upside than a real zero-cost collar allows.",
        "Solver check reprices the zero-cost call in Excel. If you change σ or r in Assumptions, re-run src/stress_test.py to refresh column Q.",
    ]
    for k, txt in enumerate(notes, n + 3):
        wc[f"A{k}"] = txt
        wc[f"A{k}"].font = f(color="555555")
    wc.freeze_panes = "B5"

    # ---------------- Stress Test ----------------
    wt = wb.create_sheet("Stress Test")
    title(wt, f"Monte Carlo stress test, {N_SCENARIOS:,} scenarios per round",
          "Each round's actual position run through lognormal settlement prices centered on the forward, at the per-round σ in Assumptions. Values from src/stress_test.py.")
    tc = ["Round", "Forward", "σ per round",
          "Expected rev, hedged ($M)", "Expected rev, unhedged ($M)",
          "Std dev, hedged ($M)", "Std dev, unhedged ($M)",
          "95% VaR, hedged ($M)", "95% VaR, unhedged ($M)", "VaR reduction",
          "95% ES, hedged ($M)", "95% ES, unhedged ($M)",
          "Worst case, hedged ($M)", "Worst case, unhedged ($M)",
          "P(bankruptcy), hedged", "P(bankruptcy), unhedged", "P(clear hurdle), hedged", "P(clear hurdle), unhedged"]
    header(wt, 4, tc, [8, 10, 10] + [13] * 15)
    keys = [("round", "0"), ("forward", USD2), ("sigma", PCT),
            ("hedged_mean_M", "0.00"), ("unhedged_mean_M", "0.00"), ("hedged_std_M", "0.00"), ("unhedged_std_M", "0.00"),
            ("hedged_var95_M", "0.00"), ("unhedged_var95_M", "0.00"), (None, PCT),
            ("hedged_es95_M", "0.00"), ("unhedged_es95_M", "0.00"), ("hedged_worst_M", "0.00"), ("unhedged_worst_M", "0.00"),
            ("hedged_p_bankrupt", PCT), ("unhedged_p_bankrupt", PCT), ("hedged_p_hurdle", PCT), ("unhedged_p_hurdle", PCT)]
    for i, (_, r) in enumerate(stress.iterrows(), 5):
        for j, (k, fmt) in enumerate(keys, 1):
            ref = f"{get_column_letter(j)}{i}"
            if k is None:
                put(wt, ref, f"=IF(I{i}=0,0,1-H{i}/I{i})", fmt, bold=True)
            else:
                v = r[k]
                put(wt, ref, None if v != v else float(v), fmt)
    last = 4 + len(stress)
    put(wt, f"A{last+1}", "Avg", bold=True, fill=TOTAL_FILL)
    for j in range(4, 19):
        col = get_column_letter(j)
        fmt = keys[j - 1][1]
        put(wt, f"{col}{last+1}", f"=IFERROR(AVERAGE({col}5:{col}{last}),\"\")", fmt, bold=True, fill=TOTAL_FILL)
    wt[f"A{last+3}"] = ("Blank probability cells = no bankruptcy line or hurdle stated in that round's memo. "
                        "Hedged revenue never falls below its floor, so 'worst case' is the floor where one exists.")
    wt[f"A{last+3}"].font = f(color="555555")
    wt.freeze_panes = "B5"

    # ---------------- Vol Sensitivity ----------------
    wv = wb.create_sheet("Vol Sensitivity")
    title(wv, "How the hedge book behaves as volatility changes",
          "Same 10,000-scenario test re-run at each σ. Rows = σ per round, columns = rounds with stated lines.")
    stated = sorted(s.dropna(subset=["bankruptcy_line"])["round"].astype(int).tolist())
    blocks = [("VaR reduction from hedging", "var_reduction_pct", None, lambda v: v / 100),
              ("P(clear hurdle): hedged minus unhedged", "hedged_p_hurdle", "unhedged_p_hurdle", None),
              ("P(bankruptcy): unhedged minus hedged", "unhedged_p_bankrupt", "hedged_p_bankrupt", None)]
    row = 4
    sigmas = sorted(sens["sigma"].unique())
    wv.column_dimensions["A"].width = 16
    for name, a, b, tf in blocks:
        rounds = sorted(sens["round"].unique()) if b is None else stated
        wv[f"A{row}"] = name
        wv[f"A{row}"].font = f(bold=True, size=11)
        header(wv, row + 1, ["σ per round"] + [f"R{r}" for r in rounds])
        for k, sg in enumerate(sigmas):
            rr = row + 2 + k
            put(wv, f"A{rr}", float(sg), PCT, BLUE)
            for j, rd in enumerate(rounds, 2):
                rec = sens[(sens["sigma"] == sg) & (sens["round"] == rd)].iloc[0]
                v = tf(rec[a]) if b is None else rec[a] - rec[b]
                put(wv, f"{get_column_letter(j)}{rr}", float(v), '+0.0%;-0.0%;"-"' if b else PCT)
        row += len(sigmas) + 4
    for j in range(2, 10):
        wv.column_dimensions[get_column_letter(j)].width = 11
    wv[f"A{row}"] = ("Read it this way: a negative hurdle number means the hedge made the investment less likely to get funded. "
                     "A positive bankruptcy number means the hedge removed that much bankruptcy probability.")
    wv[f"A{row}"].font = f(color="555555")

    # ---------------- Summary (first tab) ----------------
    wsum = wb.create_sheet("Summary", 0)
    title(wsum, "Commodity Hedging Simulation: audit report",
          "Karthik Bommisetty | UIC MBA Finance + MS Business Analytics | Fall 2025 hedging game, 7 settled rounds")
    header(wsum, 4, ["What happened in the game", "Value", "Source"], [46, 16, 40])
    srows = [
        ("Total revenue, hedged", "=Settlement!L12", MUSD, "Settlement"),
        ("Total revenue, unhedged", "=Settlement!I12", MUSD, "Settlement"),
        ("Cost of hedging", "=Settlement!I14", MUSD, "Settlement"),
        ("Cost as % of unhedged revenue", "=Settlement!I15", PCT, "Settlement"),
        ("Revenue volatility reduction", "=Settlement!I18", PCT, "Settlement"),
        ("Rounds below bankruptcy line, hedged / unhedged", '=Settlement!I19&" / "&Settlement!I20', None, "Settlement"),
        ("Rounds clearing the investment hurdle, hedged", "=Settlement!I21", "0", "Settlement"),
        ("All rounds match my memos", '=IF(COUNTIF(Settlement!O5:O11,"Yes")=7,"Yes, 7 of 7","CHECK")', None, "Settlement"),
    ]
    for i, (lab, val, fmt, src) in enumerate(srows, 5):
        put(wsum, f"A{i}", lab)
        put(wsum, f"B{i}", val, fmt, GREEN)
        put(wsum, f"C{i}", src).font = f(color="555555")
    header(wsum, 14, ["What the stress test says (before each trade)", "Value", "Source"])
    trows = [
        ("Volatility per round, estimated from game prices", "=Assumptions!B6", PCT, "Assumptions"),
        ("Average 95% VaR reduction across rounds", f"='Stress Test'!J{last+1}", PCT, "Stress Test"),
        ("Avg P(bankruptcy), unhedged", f"='Stress Test'!P{last+1}", PCT, "Stress Test, rounds with a stated line"),
        ("Avg P(bankruptcy), hedged", f"='Stress Test'!O{last+1}", PCT, "Stress Test, rounds with a stated line"),
        ("Avg P(clear hurdle), unhedged", f"='Stress Test'!R{last+1}", PCT, "Stress Test, rounds with a stated hurdle"),
        ("Avg P(clear hurdle), hedged", f"='Stress Test'!Q{last+1}", PCT, "Stress Test, rounds with a stated hurdle"),
        ("Game's edge on my 3 collars vs market pricing", f"=SUM('Collar Pricing'!P5:P{n})", MUSD, "Collar Pricing"),
    ]
    for i, (lab, val, fmt, src) in enumerate(trows, 15):
        put(wsum, f"A{i}", lab)
        put(wsum, f"B{i}", val, fmt, GREEN)
        put(wsum, f"C{i}", src).font = f(color="555555")
    wsum["A24"] = "Takeaways"
    wsum["A24"].font = f(bold=True, size=11)
    for k, txt in enumerate([
        "At the volatility this game actually had, staying unhedged carried a real chance of hitting the bankruptcy line. My hedges took that close to zero.",
        "The price was the investment: in Rounds 2 to 5 the hedges lowered the chance of funding the project. Only Round 7, with strikes near the hurdle price, raised it.",
        "Strike placement drove the outcome more than hedge ratio. Set the collar call at or above the hurdle price, so the upside given away is upside the firm doesn't need.",
    ], 25):
        wsum[f"A{k}"] = txt
        wsum[f"A{k}"].font = f()
    wsum["A29"] = "Colors: blue = input, black = formula, green = link to another sheet, yellow = key assumption."
    wsum["A29"].font = f(color="555555", italic=True)

    for rr in list(range(5, 13)) + list(range(15, 22)):
        wsum[f"B{rr}"].alignment = Alignment(horizontal="right")
    for w in wb.worksheets:
        w.sheet_view.showGridLines = False
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    print("saved", build())
